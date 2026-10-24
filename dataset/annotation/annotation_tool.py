"""Human annotation of a stratified test sample (thesis revision, Comment 3).

Three steps, run from the repository root:

  python dataset/annotation/annotation_tool.py sample     # 1. draw the sample, write one sheet per annotator
  python dataset/annotation/annotation_tool.py merge      # 2. read both sheets, kappa, write the adjudication sheet
  python dataset/annotation/annotation_tool.py finalize   # 3. read the adjudication sheet, write annotations.csv

Sheets with comment text are written to data/annotation_work/ (not in git: the sources may not be redistributed).
The result dataset/annotation/annotations.csv holds only IDs and labels, like dataset/test_label_check.csv,
and is read by notebook cell 10.4.
"""
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sklearn.metrics import cohen_kappa_score

ROOT = Path(__file__).resolve().parents[2]
DATASET = ROOT / "data" / "bangla_forensics_dataset_v1.csv.gz"
GUIDELINE = ROOT / "dataset" / "guideline_v1.md"
WORK = ROOT / "data" / "annotation_work"
OUT = ROOT / "dataset" / "annotation"

LABELS = ["normal", "offensive", "cyberbully", "hate_speech", "violence"]
CHOICES = LABELS + ["unclear"]
PER_CLASS = 80          # 5 classes x 80 = 400 test comments
SAMPLE_SEED = 2026


def draw_sample():
    """80 test comments per (source) class, spread over the sources in proportion to their size."""
    data = pd.read_csv(DATASET, encoding="utf-8-sig")
    test = data[(data["status"] == "kept") & (data["split"] == "test")]
    parts = []
    for label in LABELS:
        rows = test[test["label"] == label]
        n = min(PER_CLASS, len(rows))
        # stratified by source inside the class: sort by source, then take every k-th row of a shuffled copy
        rows = rows.sample(frac=1, random_state=SAMPLE_SEED).sort_values("source", kind="stable")
        step = len(rows) / n
        parts.append(rows.iloc[[int(i * step) for i in range(n)]])
    sample = pd.concat(parts).sample(frac=1, random_state=SAMPLE_SEED).reset_index(drop=True)
    sample.insert(0, "row", range(1, len(sample) + 1))
    return sample


def write_sheet(path, frame, label_columns, title, extra_columns=()):
    """Excel sheet: one row per comment, a drop-down for every label column, and the guideline on sheet 2."""
    columns = ["row", "id", "text", *extra_columns, *label_columns, "note"]
    frame = frame.reindex(columns=columns)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        frame.to_excel(writer, sheet_name="label", index=False)
        guideline = GUIDELINE.read_text(encoding="utf-8").split("## How the source labels")[0]
        pd.DataFrame({"guideline_v1": guideline.splitlines()}).to_excel(writer, sheet_name="guideline", index=False)
    wb = load_workbook(path)
    ws = wb["label"]
    ws.freeze_panes = "D2"
    widths = {"row": 6, "id": 20, "text": 90, "note": 30}
    for i, name in enumerate(columns, 1):
        letter = ws.cell(1, i).column_letter
        ws.column_dimensions[letter].width = widths.get(name, 16)
        ws.cell(1, i).font = Font(bold=True)
        if name in label_columns:
            dv = DataValidation(type="list", formula1='"' + ",".join(CHOICES) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{len(frame) + 1}")
            for r in range(2, len(frame) + 2):
                ws.cell(r, i).fill = PatternFill("solid", fgColor="FFF2CC")
    text_col = columns.index("text") + 1
    for r in range(2, len(frame) + 2):
        ws.cell(r, text_col).alignment = Alignment(wrap_text=True, vertical="top")
    ws.title = "label"
    ws.sheet_properties.tabColor = "1F77B4"
    wb["guideline"].column_dimensions["A"].width = 160
    wb.properties.title = title
    wb.save(path)


def read_labels(path, column):
    sheet = pd.read_excel(path, sheet_name="label", dtype=str)
    sheet[column] = sheet[column].str.strip().str.lower()
    missing = sheet[~sheet[column].isin(CHOICES)]
    if len(missing):
        raise SystemExit(f"{path.name}: {len(missing)} rows without a valid label in '{column}' "
                         f"(rows {missing['row'].head(10).tolist()} ...). Fill them in and run again.")
    return sheet[["id", column]]


def cmd_sample():
    WORK.mkdir(parents=True, exist_ok=True)
    OUT.mkdir(parents=True, exist_ok=True)
    sample = draw_sample()
    # IDs + the source label (for the notebook); the annotators never see the source label
    sample[["row", "id", "source", "label"]].rename(columns={"label": "source_label"}).to_csv(
        OUT / "sample_ids.csv", index=False)
    blind = sample[["row", "id", "text"]]
    for k in (1, 2):
        path = WORK / f"annotator_{k}.xlsx"
        if path.exists():
            raise SystemExit(f"{path} already exists - delete it first if you really want a new sheet.")
        write_sheet(path, blind, ["label"], f"Annotator {k}")
        print("written:", path)
    print(f"{len(sample)} comments:", sample["label"].value_counts().reindex(LABELS).to_dict())
    print("Give annotator_1.xlsx and annotator_2.xlsx to two different people. They must not talk about the labels.")


def cmd_merge():
    a1 = read_labels(WORK / "annotator_1.xlsx", "label").rename(columns={"label": "annotator_1"})
    a2 = read_labels(WORK / "annotator_2.xlsx", "label").rename(columns={"label": "annotator_2"})
    sample = pd.read_csv(OUT / "sample_ids.csv")
    merged = sample.merge(a1, on="id").merge(a2, on="id")
    both = merged[merged["annotator_1"].isin(LABELS) & merged["annotator_2"].isin(LABELS)]
    kappa = cohen_kappa_score(both["annotator_1"], both["annotator_2"], labels=LABELS)
    agree = (both["annotator_1"] == both["annotator_2"]).mean()
    print(f"Rows labelled by both: {len(both)} | raw agreement {agree:.1%} | Cohen's kappa {kappa:.3f}")
    print(pd.crosstab(both["annotator_1"], both["annotator_2"]))

    text = pd.read_csv(DATASET, encoding="utf-8-sig", usecols=["id", "text"])
    disputed = merged[merged["annotator_1"] != merged["annotator_2"]].merge(text, on="id")
    disputed["adjudicated"] = None
    write_sheet(WORK / "adjudication.xlsx", disputed, ["adjudicated"], "Adjudication",
                extra_columns=("annotator_1", "annotator_2"))
    merged.to_csv(WORK / "merged_annotations.csv", index=False)
    print(f"{len(disputed)} disagreements -> {WORK / 'adjudication.xlsx'} (a third person, or both annotators "
          "together, decides each one with the guideline)")


def cmd_finalize():
    merged = pd.read_csv(WORK / "merged_annotations.csv")
    adj = WORK / "adjudication.xlsx"
    decided = read_labels(adj, "adjudicated") if adj.exists() else pd.DataFrame(columns=["id", "adjudicated"])
    merged = merged.merge(decided, on="id", how="left")
    same = merged["annotator_1"] == merged["annotator_2"]
    merged.loc[same, "adjudicated"] = merged.loc[same, "annotator_1"]
    if merged["adjudicated"].isna().any():
        raise SystemExit(f"{merged['adjudicated'].isna().sum()} disagreements are not adjudicated yet.")
    cols = ["id", "source", "source_label", "annotator_1", "annotator_2", "adjudicated"]
    merged[cols].to_csv(OUT / "annotations.csv", index=False)
    print("written:", OUT / "annotations.csv", "- commit it, then run notebook cell 10.4")


if __name__ == "__main__":
    commands = {"sample": cmd_sample, "merge": cmd_merge, "finalize": cmd_finalize}
    if len(sys.argv) != 2 or sys.argv[1] not in commands:
        raise SystemExit(__doc__)
    commands[sys.argv[1]]()
